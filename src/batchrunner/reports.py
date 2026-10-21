"""Write the run's outputs: CSV, Excel, ticket comment, summary and audit log.

Outputs go to a fresh folder per run (``out/<run_id>/``). Overwriting a shared
folder would risk the worst kind of stale file: a ``closing_comment.txt`` left
over from an earlier live run sitting next to today's dry-run results.

The CSV/Excel reports carry full numbers because they are the hand-over to the
support agent who closes the tickets. The audit log and run log only ever
contain masked numbers. None of them ever contains the API token.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from batchrunner.client import AttemptRecord
from batchrunner.numbers import ParsedInput
from batchrunner.runner import Decision, NumberResult, RunReport

RESULT_COLUMNS = [
    "msisdn",
    "plan_code",
    "plan_category",
    "decision",
    "reason",
    "plan_http_status",
    "plan_attempts",
    "refresh_http_status",
    "refresh_attempts",
    "started_at",
    "finished_at",
]

DECISION_FILLS = {
    Decision.REFRESHED: "C6EFCE",
    Decision.WOULD_REFRESH: "DDEBF7",
    Decision.FLAGGED_MANUAL: "FFEB9C",
    Decision.SKIPPED_INELIGIBLE: "EDEDED",
    Decision.FAILED: "FFC7CE",
}

CLOSING_COMMENT_TEMPLATE = """Hello,

The subscription benefits for the following number(s) have been refreshed:

{numbers}

Please ask the customer to log out of the app and log back in. The benefits
appear after signing in again. If anything is still missing after that,
reply to this ticket and we will take another look.

Refreshed {count} number(s) on {date} (run {run_id}).
"""


@dataclass(frozen=True)
class OutputPaths:
    folder: Path

    @property
    def csv(self) -> Path:
        return self.folder / "results.csv"

    @property
    def excel(self) -> Path:
        return self.folder / "results.xlsx"

    @property
    def closing_comment(self) -> Path:
        return self.folder / "closing_comment.txt"

    @property
    def invalid(self) -> Path:
        return self.folder / "invalid_lines.csv"

    @property
    def unprocessed(self) -> Path:
        return self.folder / "unprocessed.txt"

    @property
    def summary(self) -> Path:
        return self.folder / "summary.json"

    @property
    def audit_log(self) -> Path:
        return self.folder / "audit_log.jsonl"

    @property
    def run_log(self) -> Path:
        return self.folder / "run.log"


class AuditLog:
    """Append-only JSONL record of every HTTP attempt, flushed line by line.

    Flushing per line means that even a run killed halfway leaves an exact
    record of which calls reached the API.
    """

    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.path = path
        self._fh = path.open("a", encoding="utf-8")

    def __call__(self, record: AttemptRecord) -> None:
        self._fh.write(json.dumps(record.to_dict()) + "\n")
        self._fh.flush()

    def close(self) -> None:
        self._fh.close()


def _row(result: NumberResult) -> dict[str, Any]:
    return {
        "msisdn": result.msisdn,
        "plan_code": result.plan_code or "",
        "plan_category": str(result.plan_category or ""),
        "decision": str(result.decision),
        "reason": result.reason,
        "plan_http_status": result.plan_status if result.plan_status is not None else "",
        "plan_attempts": result.plan_attempts,
        "refresh_http_status": result.refresh_status if result.refresh_status is not None else "",
        "refresh_attempts": result.refresh_attempts,
        "started_at": result.started_at.isoformat(timespec="seconds") if result.started_at else "",
        "finished_at": result.finished_at.isoformat(timespec="seconds")
        if result.finished_at
        else "",
    }


def write_csv(report: RunReport, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=RESULT_COLUMNS)
        writer.writeheader()
        for result in report.results:
            writer.writerow(_row(result))


def write_invalid(parsed: ParsedInput, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["line_no", "content", "reason"])
        for item in parsed.invalid:
            writer.writerow([item.line_no, item.raw, item.reason])


def write_unprocessed(report: RunReport, path: Path) -> None:
    """Numbers the run never reached, in a format ``run`` accepts as input."""
    header = f"# Not processed: {report.stopped_reason}\n"
    path.write_text(header + "\n".join(report.unprocessed) + "\n", encoding="utf-8")


def closing_comment(report: RunReport, run_id: str) -> str | None:
    """Ready-to-paste ticket comment for refreshed numbers, or None if there are none."""
    refreshed = [r.msisdn for r in report.results if r.decision == Decision.REFRESHED]
    if not refreshed:
        return None
    when = report.finished_at or report.started_at
    return CLOSING_COMMENT_TEMPLATE.format(
        numbers="\n".join(refreshed),
        count=len(refreshed),
        date=when.date().isoformat(),
        run_id=run_id,
    )


def build_summary(report: RunReport, parsed: ParsedInput, run_id: str) -> dict[str, Any]:
    steps = sum(1 for r in report.results if r.plan_attempts) + sum(
        1 for r in report.results if r.refresh_attempts
    )
    attempts = sum(r.plan_attempts + r.refresh_attempts for r in report.results)
    return {
        "run_id": run_id,
        "mode": "dry-run" if report.dry_run else "live",
        "started_at": report.started_at.isoformat(timespec="seconds"),
        "finished_at": report.finished_at.isoformat(timespec="seconds")
        if report.finished_at
        else None,
        "input": {
            "lines": parsed.total_lines,
            "blank_or_comment": parsed.ignored_lines,
            "invalid": len(parsed.invalid),
            "duplicates": len(parsed.duplicates),
            "valid_unique": len(parsed.numbers),
        },
        "decisions": {str(d): report.count(d) for d in Decision},
        "processed": len(report.results),
        "unprocessed": len(report.unprocessed),
        "http_attempts": attempts,
        "retries": attempts - steps,
        "write_attempts": sum(r.refresh_attempts for r in report.results),
        "stopped_reason": report.stopped_reason,
    }


def format_summary(summary: dict[str, Any]) -> str:
    """Human-readable multi-line summary for the console."""
    inp = summary["input"]
    lines = [
        f"Run {summary['run_id']} ({summary['mode']})",
        f"  input lines: {inp['lines']}  valid unique: {inp['valid_unique']}  "
        f"duplicates: {inp['duplicates']}  invalid: {inp['invalid']}  "
        f"blank/comment: {inp['blank_or_comment']}",
    ]
    for decision, count in summary["decisions"].items():
        if count or decision != Decision.WOULD_REFRESH or summary["mode"] == "dry-run":
            lines.append(f"  {decision:<20} {count}")
    lines.append(
        f"  HTTP attempts: {summary['http_attempts']}  retries: {summary['retries']}  "
        f"write attempts: {summary['write_attempts']}"
    )
    if summary["stopped_reason"]:
        lines.append(f"  STOPPED EARLY: {summary['stopped_reason']}")
        lines.append(f"  unprocessed numbers: {summary['unprocessed']}")
    return "\n".join(lines)


def write_excel(
    report: RunReport, parsed: ParsedInput, summary: dict[str, Any], path: Path
) -> None:
    """Formatted workbook: results (colour-coded), summary and invalid input."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(RESULT_COLUMNS)
    for result in report.results:
        row = _row(result)
        ws.append([row[c] for c in RESULT_COLUMNS])
        fill = PatternFill("solid", fgColor=DECISION_FILLS[result.decision])
        ws.cell(row=ws.max_row, column=RESULT_COLUMNS.index("decision") + 1).fill = fill
    _style_table(ws)

    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["metric", "value"])
    flat: list[tuple[str, Any]] = [
        ("run_id", summary["run_id"]),
        ("mode", summary["mode"]),
        ("started_at", summary["started_at"]),
        ("finished_at", summary["finished_at"]),
    ]
    flat += [(f"input.{k}", v) for k, v in summary["input"].items()]
    flat += [(f"decision.{k}", v) for k, v in summary["decisions"].items()]
    flat += [
        ("http_attempts", summary["http_attempts"]),
        ("retries", summary["retries"]),
        ("write_attempts", summary["write_attempts"]),
        ("unprocessed", summary["unprocessed"]),
        ("stopped_reason", summary["stopped_reason"] or ""),
    ]
    for key, value in flat:
        ws_sum.append([key, value])
    _style_table(ws_sum)

    if parsed.invalid:
        ws_inv = wb.create_sheet("Invalid input")
        ws_inv.append(["line_no", "content", "reason"])
        for item in parsed.invalid:
            ws_inv.append([item.line_no, item.raw, item.reason])
        _style_table(ws_inv)
    wb.save(path)


def _style_table(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, column in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 60)


def write_all(
    report: RunReport, parsed: ParsedInput, paths: OutputPaths, run_id: str
) -> dict[str, Any]:
    """Write every output file for a finished run and return the summary."""
    paths.folder.mkdir(parents=True, exist_ok=True)
    summary = build_summary(report, parsed, run_id)
    write_csv(report, paths.csv)
    write_excel(report, parsed, summary, paths.excel)
    paths.summary.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    if parsed.invalid:
        write_invalid(parsed, paths.invalid)
    if report.unprocessed:
        write_unprocessed(report, paths.unprocessed)
    comment = None if report.dry_run else closing_comment(report, run_id)
    if comment is not None:
        paths.closing_comment.write_text(comment, encoding="utf-8")
    return summary
