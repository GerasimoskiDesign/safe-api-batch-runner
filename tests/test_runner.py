"""End-to-end runs against the mock API over real HTTP."""

from __future__ import annotations

import csv
import json
import logging
import threading
from pathlib import Path

import httpx
from openpyxl import load_workbook

from batchrunner.pipeline import EXIT_FAILURES, EXIT_OK, EXIT_REFUSED, EXIT_STOPPED, execute
from batchrunner.runner import Decision, RunOptions
from tests.conftest import (
    AUTO,
    AUTO_2,
    FAST,
    INELIGIBLE,
    MANUAL,
    NOT_FOUND,
    TOKEN,
    UNKNOWN_PLAN,
)


def run(mock, rules, input_path, tmp_path, run_id="test-run", **options):
    return execute(
        input_path,
        mock.api,
        rules,
        RunOptions(**options),
        FAST,
        tmp_path / "out",
        run_id=run_id,
    )


def decisions(outcome) -> dict[str, str]:
    return {r.msisdn: str(r.decision) for r in outcome.report.results}


def all_output_text(folder: Path) -> str:
    """Every output file of a run as text (xlsx is a zip, so read its cells)."""
    chunks = []
    for path in folder.iterdir():
        if path.suffix == ".xlsx":
            wb = load_workbook(path)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    chunks.append(" ".join(str(v) for v in row if v is not None))
        else:
            chunks.append(path.read_text(encoding="utf-8"))
    return "\n".join(chunks)


# --- decisions ---------------------------------------------------------------


def test_happy_path_refreshes_eligible_numbers(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input(AUTO, AUTO_2), tmp_path)
    assert outcome.exit_code == EXIT_OK
    assert decisions(outcome) == {AUTO: "REFRESHED", AUTO_2: "REFRESHED"}
    stats = mock.stats()
    assert stats["plan_calls"] == 2 and stats["writes_applied"] == 2
    assert dict(mock.state.refreshed) == {AUTO: 1, AUTO_2: 1}


def test_ineligible_and_unknown_plans_are_skipped_without_writes(
    start_mock, rules, write_input, tmp_path
):
    mock = start_mock()
    outcome = run(mock, rules, write_input(INELIGIBLE, UNKNOWN_PLAN), tmp_path)
    assert decisions(outcome) == {
        INELIGIBLE: "SKIPPED_INELIGIBLE",
        UNKNOWN_PLAN: "SKIPPED_INELIGIBLE",
    }
    reasons = {r.msisdn: r.reason for r in outcome.report.results}
    assert "not in config" in reasons[UNKNOWN_PLAN]
    assert mock.stats()["write_calls"] == 0


def test_manual_only_plan_is_flagged_not_refreshed(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input(MANUAL), tmp_path)
    assert decisions(outcome) == {MANUAL: "FLAGGED_MANUAL"}
    assert mock.stats()["write_calls"] == 0


def test_mixed_batch_records_status_and_attempts(start_mock, rules, write_input, tmp_path):
    mock = start_mock(scripted_faults={f"refresh:{AUTO}": ["5xx", "ok"]})
    outcome = run(mock, rules, write_input(AUTO, INELIGIBLE, MANUAL, NOT_FOUND), tmp_path)
    assert outcome.exit_code == EXIT_FAILURES  # the 404 number
    by_number = {r.msisdn: r for r in outcome.report.results}
    assert by_number[AUTO].decision is Decision.REFRESHED
    assert (by_number[AUTO].plan_status, by_number[AUTO].refresh_status) == (200, 200)
    assert by_number[AUTO].refresh_attempts == 2
    assert by_number[NOT_FOUND].decision is Decision.FAILED
    assert by_number[NOT_FOUND].plan_status == 404
    with outcome.paths.csv.open(encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert [r["msisdn"] for r in rows] == [AUTO, INELIGIBLE, MANUAL, NOT_FOUND]
    assert rows[0]["refresh_attempts"] == "2"
    assert rows[1]["refresh_http_status"] == ""


# --- dry run -----------------------------------------------------------------


def test_dry_run_makes_zero_write_calls(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(
        mock, rules, write_input(AUTO, AUTO_2, MANUAL, INELIGIBLE), tmp_path, dry_run=True
    )
    stats = mock.stats()
    assert stats["write_calls"] == 0 and stats["writes_applied"] == 0
    assert stats["plan_calls"] == 4
    assert decisions(outcome)[AUTO] == "WOULD_REFRESH"
    assert not outcome.paths.closing_comment.exists()
    assert outcome.summary["mode"] == "dry-run"


# --- cap ---------------------------------------------------------------------


def numbers(count: int) -> list[str]:
    return [f"+9997100{i:04d}" for i in range(count)]


def test_cap_refuses_oversized_batch_before_any_call(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input(*numbers(51)), tmp_path)
    assert outcome.exit_code == EXIT_REFUSED
    assert "cap" in outcome.error
    assert mock.stats()["total_calls"] == 0
    assert not (tmp_path / "out").exists()


def test_cap_counts_unique_valid_numbers_only(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    lines = numbers(50) + numbers(5) + ["garbage", "# comment"]
    outcome = run(mock, rules, write_input(*lines), tmp_path, dry_run=True)
    assert outcome.exit_code in (EXIT_OK, EXIT_FAILURES)
    assert mock.stats()["plan_calls"] == 50


def test_allow_over_cap_is_an_explicit_opt_in(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(
        mock,
        rules,
        write_input(*numbers(4)),
        tmp_path,
        dry_run=True,
        max_numbers=3,
        allow_over_cap=True,
    )
    assert outcome.exit_code != EXIT_REFUSED
    assert mock.stats()["plan_calls"] == 4


# --- circuit breaker & auth --------------------------------------------------


def test_circuit_breaker_stops_run_and_lists_leftovers(start_mock, rules, write_input, tmp_path):
    mock = start_mock(error_rate=1.0)
    batch = numbers(10)
    outcome = run(mock, rules, write_input(*batch), tmp_path, breaker_threshold=3)
    assert outcome.exit_code == EXIT_STOPPED
    assert "circuit breaker" in outcome.report.stopped_reason
    assert len(outcome.report.results) == 3
    assert outcome.report.unprocessed == batch[3:]
    # 3 numbers x 3 attempts, then nothing more reaches the backend.
    assert mock.stats()["total_calls"] == 9
    leftovers = outcome.paths.unprocessed.read_text(encoding="utf-8").splitlines()
    assert leftovers[0].startswith("# Not processed")
    assert leftovers[1:] == batch[3:]


def test_a_success_resets_the_breaker(start_mock, rules, write_input, tmp_path):
    a, b, c, d = numbers(4)
    faults = {f"plan:{n}": ["5xx"] * 3 for n in (a, b, d)}
    mock = start_mock(scripted_faults=faults)
    outcome = run(mock, rules, write_input(a, b, c, d), tmp_path, breaker_threshold=3)
    # a, b fail; c succeeds and resets the count; d fails -> only 1 in a row.
    assert outcome.report.stopped_reason is None
    assert len(outcome.report.results) == 4


def test_not_found_numbers_do_not_trip_the_breaker(start_mock, rules, write_input, tmp_path):
    missing = numbers(5)
    mock = start_mock(plan_overrides=dict.fromkeys(missing))
    outcome = run(mock, rules, write_input(*missing), tmp_path, breaker_threshold=2)
    assert outcome.report.stopped_reason is None
    assert outcome.report.count(Decision.FAILED) == 5
    assert mock.stats()["total_calls"] == 5  # 404s are not retried


def test_bad_token_stops_after_first_call(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = execute(
        write_input(AUTO, AUTO_2, MANUAL),
        mock.api.__class__(base_url=mock.base_url, token="expired-credential"),
        rules,
        RunOptions(),
        FAST,
        tmp_path / "out",
        run_id="bad-token",
    )
    assert outcome.exit_code == EXIT_STOPPED
    assert "credentials" in outcome.report.stopped_reason
    assert mock.stats()["total_calls"] == 1


# --- strictly sequential -----------------------------------------------------


def test_server_never_sees_more_than_one_request_in_flight(
    start_mock, rules, write_input, tmp_path
):
    batch = numbers(12) + [AUTO, AUTO_2]
    faults = {f"plan:{batch[2]}": ["slow", "ok"], f"refresh:{AUTO}": ["5xx", "ok"]}
    mock = start_mock(error_rate=0.1, scripted_faults=faults, slow_seconds=1.0)
    run(mock, rules, write_input(*batch), tmp_path, breaker_threshold=5)
    stats = mock.stats()
    assert stats["total_calls"] >= len(batch)
    assert stats["dropped_by_client"] >= 1  # a timeout really happened
    assert stats["max_concurrency"] == 1


def test_concurrency_meter_does_detect_parallel_calls(start_mock):
    """Control experiment: prove /admin/stats would catch a parallel client."""
    mock = start_mock(scripted_faults={f"plan:{AUTO}": ["slow"] * 4}, slow_seconds=0.5)
    headers = {"Authorization": f"Bearer {TOKEN}"}
    url = f"{mock.base_url}/subscribers/%2B99970000001/plan"
    threads = [
        threading.Thread(target=httpx.get, args=(url,), kwargs={"headers": headers})
        for _ in range(4)
    ]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert mock.stats()["max_concurrency"] > 1


# --- secrets and masking -----------------------------------------------------


def test_token_never_appears_in_logs_or_outputs(start_mock, rules, write_input, tmp_path, caplog):
    caplog.set_level(logging.DEBUG)
    mock = start_mock(
        scripted_faults={f"plan:{AUTO}": ["5xx", "ok"], f"plan:{AUTO_2}": ["slow", "ok"]}
    )
    outcome = run(mock, rules, write_input(AUTO, AUTO_2, MANUAL, NOT_FOUND, "bad line"), tmp_path)
    files = list(outcome.paths.folder.iterdir())
    assert {p.name for p in files} >= {
        "results.csv",
        "results.xlsx",
        "summary.json",
        "audit_log.jsonl",
        "run.log",
        "closing_comment.txt",
        "invalid_lines.csv",
    }
    assert TOKEN not in caplog.text  # raw records, before any redaction
    assert TOKEN not in all_output_text(outcome.paths.folder)


def test_logs_and_audit_log_only_contain_masked_numbers(
    start_mock, rules, write_input, tmp_path, caplog
):
    caplog.set_level(logging.DEBUG)
    mock = start_mock(scripted_faults={f"plan:{AUTO}": ["5xx", "ok"]})
    outcome = run(mock, rules, write_input(AUTO, MANUAL, NOT_FOUND), tmp_path)
    log_text = outcome.paths.run_log.read_text(encoding="utf-8")
    audit_text = outcome.paths.audit_log.read_text(encoding="utf-8")
    for number in (AUTO, MANUAL, NOT_FOUND):
        digits = number[1:]
        assert digits not in caplog.text
        assert digits not in log_text
        assert digits not in audit_text
    assert "+********001" in log_text
    audit = [json.loads(line) for line in audit_text.splitlines()]
    assert [a["status_code"] for a in audit if a["msisdn"] == "+********001"][:2] == [503, 200]


# --- outputs -----------------------------------------------------------------


def test_closing_comment_lists_only_refreshed_numbers(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input(AUTO, MANUAL, INELIGIBLE, AUTO_2), tmp_path)
    comment = outcome.paths.closing_comment.read_text(encoding="utf-8")
    assert AUTO in comment and AUTO_2 in comment
    assert MANUAL not in comment and INELIGIBLE not in comment
    assert "log out" in comment and "log back in" in comment


def test_no_closing_comment_when_nothing_was_refreshed(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input(MANUAL, INELIGIBLE), tmp_path)
    assert not outcome.paths.closing_comment.exists()


def test_excel_report_has_formatted_sheets(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input(AUTO, MANUAL, "nonsense"), tmp_path)
    wb = load_workbook(outcome.paths.excel)
    assert wb.sheetnames == ["Results", "Summary", "Invalid input"]
    ws = wb["Results"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold
    header = [c.value for c in ws[1]]
    assert header[:4] == ["msisdn", "plan_code", "plan_category", "decision"]
    assert ws.max_row == 3


def test_existing_output_folder_is_never_overwritten(start_mock, rules, write_input, tmp_path):
    import pytest

    from batchrunner.config import ConfigError

    mock = start_mock()
    path = write_input(AUTO)
    run(mock, rules, path, tmp_path, run_id="same")
    with pytest.raises(ConfigError, match="already exists"):
        run(mock, rules, path, tmp_path, run_id="same")


def test_no_valid_numbers_refuses_to_start(start_mock, rules, write_input, tmp_path):
    mock = start_mock()
    outcome = run(mock, rules, write_input("# only a comment", "junk"), tmp_path)
    assert outcome.exit_code == EXIT_REFUSED
    assert mock.stats()["total_calls"] == 0
